import io
import os
from pathlib import Path
from datetime import timedelta
from django.test import TestCase, override_settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.utils import timezone
from PIL import Image
from pypdf import PdfWriter
from docx import Document
from . import test_community
from .models import TeamOpening, TeamApplication, PersonalMessage, MessageUpload, Team, Workspace, TeamMembership, Sticker, AccountPreferences
from .tenancy import scope
from unittest.mock import patch


class ExperienceRepair(TestCase):
    setUp=test_community.CommunityTests.setUp
    def capture(self,name,response):
        folder=os.environ.get('WORKBENCH_CAPTURE_UI')
        if folder:Path(folder,name+'.html').write_bytes(response.content)

    def upload(self,name,content):
        row=PersonalMessage.objects.create(sender=self.owner,recipient=self.worker,body='',relation_verified=True)
        return MessageUpload.objects.create(personal=row,original_name=name,file=SimpleUploadedFile(name,content))

    def test_pdf_preview_and_page_render(self):
        writer=PdfWriter();writer.add_blank_page(width=612,height=792);buffer=io.BytesIO();writer.write(buffer)
        item=self.upload('方案.pdf',buffer.getvalue());url=reverse('file_preview',args=['upload',item.pk])
        page=self.client.get(url);self.assertContains(page,'file-document-page');self.capture('repair-file-pdf',page)
        image=self.client.get(reverse('file_preview_content',args=['upload',item.pk]),{'page':1});self.assertEqual(image.status_code,200);self.assertEqual(image['Content-Type'],'image/png')
        self.assertEqual(self.client.get(reverse('file_preview_content',args=['upload',item.pk]),{'page':'0'}).status_code,404)
        self.client.force_login(self.outside);self.assertEqual(self.client.get(url).status_code,404)

    def test_document_text_is_escaped_and_tables_are_previewed(self):
        doc=Document();doc.add_paragraph('<script>bad()</script>');table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='项目';table.cell(0,1).text='正常内容';buffer=io.BytesIO();doc.save(buffer)
        item=self.upload('SOP2.docx',buffer.getvalue());response=self.client.get(reverse('file_preview',args=['upload',item.pk]));self.assertContains(response,'&lt;script&gt;');self.assertNotContains(response,'<script>bad()');self.assertContains(response,'正常内容');self.capture('repair-file-docx',response)

    def test_preview_rechecks_hidden_and_withdrawn_files(self):
        item=self.upload('read.txt',b'local file');url=reverse('file_preview',args=['upload',item.pk]);self.assertEqual(self.client.get(url).status_code,200)
        item.personal.hidden_by.add(self.owner);self.assertEqual(self.client.get(url).status_code,403)
        item.personal.hidden_by.remove(self.owner);PersonalMessage.objects.filter(pk=item.personal_id).update(withdrawn_at=timezone.now());self.assertEqual(self.client.get(url).status_code,403)

    def test_message_files_have_metadata_and_preview_link(self):
        item=self.upload('SOP2.docx',b'12345');response=self.client.get(reverse('personal_chat',args=[self.worker.pk]),HTTP_ACCEPT='application/json');file=response.json()['messages'][0]['files'][0]
        self.assertEqual(file['size'],5);self.assertEqual(file['type'],'DOCX');self.assertEqual(file['preview_url'],reverse('file_preview',args=['upload',item.pk]))
        response=self.client.get(reverse('personal_chat',args=[self.worker.pk]));self.assertContains(response,'file-card-copy');self.assertContains(response,'data-download-url')

    def test_poll_updates_only_messages_already_known_and_does_not_starve_new_messages(self):
        old=PersonalMessage.objects.create(sender=self.owner,recipient=self.worker,body='old',relation_verified=True,withdrawn_at=timezone.now())
        newest=PersonalMessage.objects.create(sender=self.worker,recipient=self.owner,body='new',relation_verified=True)
        url=reverse('personal_chat',args=[self.worker.pk])
        response=self.client.get(url,{'after':old.pk},HTTP_ACCEPT='application/json');self.assertEqual([r['id'] for r in response.json()['messages']],[newest.pk])
        response=self.client.get(url,{'after':newest.pk,'known':old.pk},HTTP_ACCEPT='application/json');self.assertEqual([r['id'] for r in response.json()['messages']],[old.pk])

    def test_no_team_account_can_upload_and_manage_stickers(self):
        self.client.force_login(self.outside);data=io.BytesIO();Image.new('RGB',(10,10),'blue').save(data,'PNG')
        response=self.client.post('/messages/stickers/',{'file':SimpleUploadedFile('贴图.png',data.getvalue(),content_type='image/png')});self.assertEqual(response.status_code,201,response.content);identifier=response.json()['id']
        self.assertEqual(self.client.get(response.json()['url']).status_code,200)
        self.assertEqual(self.client.post('/messages/stickers/',{'action':'move','ids':[identifier]}).status_code,200)
        self.assertEqual(AccountPreferences.objects.get(user=self.outside).sticker_order,[identifier])
        page=self.client.get('/app/pages/',{'view':'stickers'});self.assertContains(page,'data-sticker-manager');self.capture('repair-stickers',page)
        self.assertEqual(self.client.post('/messages/stickers/',{'action':'bulk-remove','ids':[identifier]}).status_code,200)
        self.assertEqual(self.client.get('/messages/stickers/').json()['stickers'],[])

    def test_sticker_bulk_actions_reject_other_accounts_favorites(self):
        with scope(self.team):sticker=Sticker.objects.create(created_by=self.owner,name='private',mime='image/png',file=SimpleUploadedFile('s.png',b'data'))
        sticker.favorites.add(self.owner);self.client.force_login(self.outside)
        self.assertEqual(self.client.post('/messages/stickers/',{'action':'bulk-remove','ids':[sticker.pk]}).status_code,403)
        self.assertTrue(sticker.favorites.filter(pk=self.owner.pk).exists())

    def test_recruitment_pause_edit_resume_close_and_full_guard(self):
        opening=TeamOpening.objects.create(team=self.team,title='职位',description='说明',planned_headcount=1)
        url=reverse('recruitment_manage')
        self.assertEqual(self.client.post(url,{'action':'pause','opening':opening.pk}).status_code,302);opening.refresh_from_db();self.assertTrue(opening.paused);self.assertFalse(opening.active)
        self.client.post(url,{'action':'edit','opening':opening.pk,'title':'新职位','description':'新说明','planned_headcount':2});opening.refresh_from_db();self.assertEqual(opening.title,'新职位');self.assertTrue(opening.paused)
        self.client.post(url,{'action':'resume','opening':opening.pk});opening.refresh_from_db();self.assertTrue(opening.active);self.assertFalse(opening.paused)
        self.client.post(url,{'action':'close','opening':opening.pk});opening.refresh_from_db();self.assertFalse(opening.active);self.assertFalse(opening.paused)
        TeamApplication.objects.create(opening=opening,applicant=self.outside,resume='简历',state='joined');TeamOpening.objects.filter(pk=opening.pk).update(planned_headcount=1)
        self.client.post(url,{'action':'resume','opening':opening.pk});opening.refresh_from_db();self.assertFalse(opening.active)
        page=self.client.get(url);self.assertContains(page,'编辑职位');self.assertContains(page,'重新开放');self.capture('repair-recruitment',page)

    def test_recruitment_management_is_owner_or_hr_only(self):
        self.client.force_login(self.worker);self.assertEqual(self.client.post(reverse('recruitment_manage'),{'action':'pause','opening':1}).status_code,403)

    def test_solo_team_disband_keeps_history_and_does_not_create_replacement_team(self):
        solo=Team.objects.create(name='封闭团队',owner=self.owner,solo_owner=self.owner,recruitment_mode='closed');space=Workspace.objects.create(kind='team',team=solo)
        TeamMembership.objects.create(team=solo,user=self.owner,role='owner')
        session=self.client.session;session['workbench-team']=solo.pk;session['message-team']=solo.pk;session['workbench-space']='team:'+str(solo.pk);session.save()
        before=Team.objects.count()
        response=self.client.post(reverse('team_disband'),{'confirm':solo.name});self.assertEqual(response.status_code,302);solo.refresh_from_db();space.refresh_from_db();self.assertFalse(solo.active);self.assertFalse(space.active)
        self.assertEqual(Team.objects.count(),before);self.assertNotIn('message-team',self.client.session)

    def test_appearance_page_is_available_without_team(self):
        self.client.force_login(self.outside);page=self.client.get('/me/appearance/');self.assertContains(page,'data-accent-choice="purple"');self.capture('repair-appearance',page)

    def test_file_favorites_remain_private_and_withdrawal_revokes_preview(self):
        item=self.upload('read.txt',b'content');url=reverse('file_favorite',args=['upload',item.pk])
        self.assertEqual(self.client.post(url,{'action':'add'}).json(),{'saved':True})
        self.assertContains(self.client.get('/me/files/'),'read.txt')
        self.client.force_login(self.outside);self.assertNotContains(self.client.get('/me/files/'),'read.txt')
        self.assertEqual(self.client.post(url,{'action':'add'}).status_code,404)
        self.client.force_login(self.owner);PersonalMessage.objects.filter(pk=item.personal_id).update(withdrawn_at=timezone.now())
        self.assertNotContains(self.client.get('/me/files/'),'read.txt')
        self.assertEqual(self.client.post(url,{'action':'remove'}).json(),{'saved':False})

    def test_legacy_word_conversion_renders_pdf_pages(self):
        from tempfile import TemporaryDirectory
        writer=PdfWriter();writer.add_blank_page(width=612,height=792)
        with TemporaryDirectory() as folder:
            target=Path(folder)/'render.pdf'
            with target.open('wb') as output:writer.write(output)
            item=self.upload('说明.doc',b'word')
            with patch('core.office_preview.pdf_path',return_value=target):
                self.assertContains(self.client.get(reverse('file_preview',args=['upload',item.pk])),'file-document-page')
                response=self.client.get(reverse('file_preview_content',args=['upload',item.pk]),{'page':1});self.assertEqual(response['Content-Type'],'image/png')

    def test_converted_solo_disband_preserves_private_ai_history(self):
        from aihub.models import AssistantConversation, AssistantJob
        from .solo_teams import account_storage
        solo=Team.objects.create(name='旧单人团队',owner=self.owner,solo_owner=self.owner,recruitment_mode='closed')
        space=Workspace.objects.create(kind='team',team=solo,owner=self.owner)
        TeamMembership.objects.create(team=solo,user=self.owner,role='owner')
        with scope(space):
            conversation=AssistantConversation.objects.create(user=self.owner,title='保留的对话')
            job=AssistantJob.objects.create(user=self.owner,conversation=conversation,state='done',billing_workspace=space)
        session=self.client.session;session['workbench-team']=solo.pk;session['workbench-space']='team:'+str(solo.pk);session.save()
        self.client.post(reverse('team_disband'),{'confirm':solo.name})
        private=account_storage(self.owner);self.assertEqual(private.kind,'personal');self.assertTrue(private.active)
        conversation.refresh_from_db();job.refresh_from_db();self.assertEqual(conversation.workspace_id,private.pk);self.assertEqual(job.workspace_id,private.pk);self.assertEqual(job.billing_workspace_id,space.pk)
        self.assertContains(self.client.get('/me/appearance/'),'data-accent-choice="purple"')

    def test_disband_waits_for_live_ai_call_and_preserves_billing(self):
        from aihub.models import AssistantJob
        space=Workspace.objects.create(kind='team',team=self.team)
        with scope(space):
            job=AssistantJob.objects.create(user=self.owner,billing_workspace=space)
        response=self.client.post(reverse('team_disband'),{'confirm':self.team.name})
        self.assertEqual(response.status_code,302)
        self.team.refresh_from_db();job.refresh_from_db()
        self.assertTrue(self.team.active);self.assertEqual(job.state,'running')
        self.assertEqual(job.billing_workspace_id,space.pk)
