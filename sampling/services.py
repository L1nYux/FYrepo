from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import re
import unicodedata
import zipfile
from collections import defaultdict
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import parse_qs, urlparse

from django.core.files.base import ContentFile
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Count
from django.utils import timezone

from .models import CandidatePaper, SamplingArtifact, SamplingRun
from .scope import tier_for, validate_scope


CONFIG_DIR = Path(__file__).resolve().parent / "config"
REGISTRY = json.loads((CONFIG_DIR / "journal_registry.json").read_text(encoding="utf-8"))
RULES = json.loads((CONFIG_DIR / "eligibility_rules.json").read_text(encoding="utf-8"))
PROTOCOL_VERSION = "sampling_protocol_v1"

ALIASES = {
    "题名": "title", "论文题名": "title", "title": "title",
    "作者": "authors", "authors": "authors", "author": "authors",
    "文献来源": "journal", "来源": "journal", "期刊": "journal",
    "journal": "journal", "source": "journal",
    "年": "year", "年份": "year", "发表时间": "year", "出版日期": "year",
    "year": "year", "pub_date": "year",
    "期": "issue", "期号": "issue", "issue": "issue",
    "卷": "volume", "volume": "volume",
    "DOI": "doi", "doi": "doi",
    "cnki_url": "url", "标题": "title", "知网链接": "url", "URL": "url", "网址": "url", "链接": "url", "detail_url": "url", "url": "url",
    "关键词": "keywords", "keywords": "keywords",
    "摘要": "abstract", "abstract": "abstract",
    "中图分类号": "clc", "CLC": "clc", "clc": "clc",
    "来源库": "database", "database": "database",
    "文献类型": "article_type", "article_type": "article_type", "kind": "article_type",
    "filename": "cnki_filename", "cnki_filename": "cnki_filename",
    "dbcode": "cnki_dbcode", "cnki_dbcode": "cnki_dbcode",
}
CANON = [
    "title", "authors", "journal", "year", "issue", "volume", "doi", "url",
    "keywords", "abstract", "clc", "database", "article_type",
    "cnki_filename", "cnki_dbcode",
]


class SamplingError(ValueError):
    pass


class SamplingShortage(SamplingError):
    def __init__(self, shortages):
        self.shortages = shortages
        detail = "；".join(
            f"{x['stratum_id']} 需要{x['requested']}，可用{x['available']}"
            for x in shortages[:12]
        )
        if len(shortages) > 12:
            detail += f"；另有 {len(shortages)-12} 个抽样格不足"
        super().__init__(f"主样本/留出样本数量不足：{detail}")


def clean(value) -> str:
    if value is None:
        return ""
    return re.sub(
        r"\s+", " ",
        unicodedata.normalize("NFKC", str(value)).replace("\xa0", " ")
    ).strip()


def norm(value) -> str:
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]+", "", clean(value).lower())


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def stable_hash(seed: str, *parts: str) -> str:
    return sha256_text("|".join([clean(seed), *[clean(p) for p in parts]]))


def parse_cnki(url: str, filename: str = "", dbcode: str = "") -> dict:
    url, filename, dbcode = clean(url), clean(filename), clean(dbcode)
    if url:
        try:
            qs = parse_qs(urlparse(url).query)
            for k in ("filename", "FileName", "fileName", "FILENAME"):
                if not filename and qs.get(k):
                    filename = clean(qs[k][0])
            for k in ("dbcode", "DBCODE", "DbCode"):
                if not dbcode and qs.get(k):
                    dbcode = clean(qs[k][0])
        except Exception:
            pass
    if dbcode and filename:
        key, method = f"{dbcode}:{filename}", "dbcode+filename"
    elif filename:
        key, method = filename, "filename"
    elif url:
        key, method = "URL:" + sha256_text(url)[:16], "url_hash"
    else:
        key, method = "", "missing"
    return {
        "cnki_dbcode": dbcode,
        "cnki_filename": filename,
        "cnki_record_key": key,
        "record_key_method": method,
    }


def normalize_doi(value: str) -> str:
    x = clean(value).lower()
    x = re.sub(r"^https?://(?:dx\.)?doi\.org/", "", x)
    x = re.sub(r"^doi\s*[:：]?\s*", "", x)
    return x.rstrip(".,;；")


def candidate_key(rec: dict) -> str:
    ids = parse_cnki(
        rec.get("url", ""), rec.get("cnki_filename", ""), rec.get("cnki_dbcode", "")
    )
    if ids["cnki_record_key"] and ids["record_key_method"] != "url_hash":
        return "cnki:" + ids["cnki_record_key"].lower()
    doi = normalize_doi(rec.get("doi", ""))
    if doi:
        return "doi:" + doi
    if clean(rec.get("url", "")):
        return "url:" + clean(rec["url"]).lower().rstrip("/")
    author = re.split(r"[;；,，、]", clean(rec.get("authors", "")))[0]
    return "meta:" + "|".join([
        norm(rec.get("title", "")),
        clean(rec.get("year", ""))[:4],
        norm(rec.get("journal", "")),
        norm(author),
    ])


def _extract_year(value):
    m = re.search(r"(?:19|20)\d{2}", clean(value))
    return int(m.group(0)) if m else None


def normalize_record(raw: dict, source_file="") -> dict:
    rec = {ALIASES.get(clean(k), clean(k)): clean(v) for k, v in raw.items()}
    out = {k: clean(rec.get(k, "")) for k in CANON}
    out["year"] = _extract_year(out["year"])
    ids = parse_cnki(out["url"], out["cnki_filename"], out["cnki_dbcode"])
    out.update(ids)
    out["candidate_key"] = candidate_key(out)
    out["source_file"] = source_file
    return out


class _HTMLTableParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tables = []
        self.table = None
        self.row = None
        self.cell = None

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if tag == "table" and self.table is None:
            self.table = []
        elif self.table is not None and tag == "tr":
            self.row = []
        elif self.row is not None and tag in ("td", "th"):
            self.cell = []

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag in ("td", "th") and self.cell is not None and self.row is not None:
            self.row.append(clean("".join(self.cell)))
            self.cell = None
        elif tag == "tr" and self.row is not None and self.table is not None:
            if any(self.row):
                self.table.append(self.row)
            self.row = None
        elif tag == "table" and self.table is not None:
            if self.table:
                self.tables.append(self.table)
            self.table = None


def _decode_bytes(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    return data.decode("utf-8", errors="replace")


def _records_from_csv(data: bytes):
    text = _decode_bytes(data)
    return list(csv.DictReader(io.StringIO(text)))


def _records_from_json(data: bytes):
    obj = json.loads(_decode_bytes(data))
    if isinstance(obj, dict):
        obj = obj.get("rows", obj.get("records", obj.get("data", obj)))
    if not isinstance(obj, list):
        raise SamplingError("JSON 必须是记录数组，或包含 rows / records / data 数组。")
    return [x for x in obj if isinstance(x, dict)]


def _records_from_xlsx(data: bytes):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SamplingError("服务器缺少 openpyxl，无法读取 XLSX。") from exc
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [clean(x) for x in rows[0]]
    return [
        {headers[i]: row[i] if i < len(row) else "" for i in range(len(headers))}
        for row in rows[1:] if any(clean(x) for x in row)
    ]


def _records_from_html_xls(data: bytes):
    parser = _HTMLTableParser()
    parser.feed(_decode_bytes(data))
    if not parser.tables:
        raise SamplingError("未在 XLS 文件中识别到 HTML 表格。")
    table = max(parser.tables, key=len)
    if len(table) < 2:
        return []
    headers = [clean(x) for x in table[0]]
    return [
        {headers[i]: row[i] if i < len(row) else "" for i in range(len(headers))}
        for row in table[1:] if any(clean(x) for x in row)
    ]


def read_metadata_bytes(filename: str, data: bytes) -> list[dict]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        rows = _records_from_csv(data)
    elif suffix == ".json":
        rows = _records_from_json(data)
    elif suffix in (".xlsx", ".xlsm"):
        rows = _records_from_xlsx(data)
    elif suffix == ".xls":
        rows = _records_from_html_xls(data)
    else:
        raise SamplingError("仅支持 CSV、JSON、XLSX/XLSM、知网 HTML-XLS。")
    normalized = [normalize_record(row, filename) for row in rows]
    seen = set()
    result = []
    for rec in normalized:
        if not rec["candidate_key"] or rec["candidate_key"] in seen:
            continue
        seen.add(rec["candidate_key"])
        result.append(rec)
    return result


def tier_lookup():
    return {journal: tier for tier, journals in REGISTRY["tiers"].items() for journal in journals}


def build_frame(run: SamplingRun) -> list[dict]:
    if run.is_frozen:
        if run.protocol.get('frame_snapshot'):
            return run.protocol['frame_snapshot']
        artifact = run.artifacts.filter(artifact_type=SamplingArtifact.TYPE_FRAME).first()
        if not artifact:
            return []
        with artifact.file.open('rb') as source:
            rows = list(csv.DictReader(io.StringIO(_decode_bytes(source.read()))))
        for row in rows:
            for field in ('start_year', 'end_year', 'target_n', 'reserve_n', 'cross_disciplinary_journal', 'frame_order'):
                row[field] = int(row[field])
        return rows
    tiers = tier_lookup()
    rows = []
    order = 1
    for journal in run.selected_journals:
        base_tier = tiers.get(journal, "")
        aliases = REGISTRY.get("journal_aliases", {}).get(journal, [journal])
        for period in run.periods:
            tier = (
                REGISTRY.get("period_tier_overrides", {})
                .get(journal, {})
                .get(period["id"], base_tier)
            )
            if tier not in run.selected_tiers:
                continue
            rows.append({
                "stratum_id": f"MAIN_{order:04d}",
                "role": "MAIN",
                "journal_family_name": journal,
                "journal_query_names": "|".join(aliases),
                "period": period["id"],
                "start_year": int(period["start_year"]),
                "end_year": int(period["end_year"]),
                "tier": tier,
                "target_n": int(run.main_n),
                "reserve_n": int(run.reserve_n),
                "cross_disciplinary_journal": 1 if journal in REGISTRY.get("cross_disciplinary", []) else 0,
                "frame_order": order,
            })
            order += 1
    if run.holdout_enabled:
        for journal in run.selected_journals:
            base_tier = tiers.get(journal, "")
            if tier_for(journal, "P3") not in run.selected_tiers:
                continue
            aliases = REGISTRY.get("journal_aliases", {}).get(journal, [journal])
            rows.append({
                "stratum_id": f"HOLDOUT_{order:04d}",
                "role": "HOLDOUT",
                "journal_family_name": journal,
                "journal_query_names": "|".join(aliases),
                "period": "HOLDOUT",
                "start_year": int(run.holdout_start),
                "end_year": int(run.holdout_end),
                "tier": (
                    REGISTRY.get("period_tier_overrides", {})
                    .get(journal, {})
                    .get("P3", base_tier)
                ),
                "target_n": int(run.holdout_n),
                "reserve_n": int(run.reserve_n),
                "cross_disciplinary_journal": 1 if journal in REGISTRY.get("cross_disciplinary", []) else 0,
                "frame_order": order,
            })
            order += 1
    return rows


def _attach_frame(run: SamplingRun, rec: dict):
    jr = norm(rec.get("journal", ""))
    yr = rec.get("year")
    for fr in build_frame(run):
        aliases = [norm(x) for x in fr["journal_query_names"].split("|") if clean(x)]
        if yr is not None and jr in aliases and fr["start_year"] <= int(yr) <= fr["end_year"]:
            return fr
    return None


def _classify(rec: dict, fr: dict | None):
    title = clean(rec.get("title"))
    clc = clean(rec.get("clc")).upper()
    if fr is None:
        return CandidatePaper.HOLD_METADATA, "未匹配当前抽样范围", CandidatePaper.EXCLUDE
    if not title or not clean(rec.get("journal")) or not rec.get("year"):
        return CandidatePaper.HOLD_METADATA, "缺题名/期刊/年份", CandidatePaper.EXCLUDE
    for term in RULES["hard_exclude_title_terms"]:
        if term in title:
            return CandidatePaper.EXCLUDED_AUTO, f"标题命中排除词：{term}", CandidatePaper.EXCLUDE
    for term in RULES["uncertain_title_terms"]:
        if term in title:
            return CandidatePaper.UNCERTAIN, f"需人工确认：{term}", CandidatePaper.PENDING
    if fr["cross_disciplinary_journal"] and not (
        clc and clc.startswith(tuple(RULES.get("law_clc_prefixes", ["D9"])))
    ):
        return CandidatePaper.UNCERTAIN, "跨学科期刊：需确认是否为法学论文", CandidatePaper.PENDING
    return CandidatePaper.ELIGIBLE_AUTO, "研究论文候选", CandidatePaper.INCLUDE


def refresh_run_status(run: SamplingRun):
    if run.status in (SamplingRun.FROZEN, SamplingRun.ARCHIVED):
        return
    unresolved = run.candidates.filter(
        eligibility_status=CandidatePaper.UNCERTAIN,
        decision=CandidatePaper.PENDING,
    ).exists()
    if unresolved:
        status = SamplingRun.REVIEW
    elif run.candidates.exists():
        status = SamplingRun.READY
    else:
        status = SamplingRun.DRAFT
    if run.status != status:
        SamplingRun.objects.filter(pk=run.pk).update(status=status, updated_at=timezone.now())
        run.status = status


@transaction.atomic
def ingest_records(run: SamplingRun, records: list[dict], source_file="") -> dict:
    run = SamplingRun.objects.select_for_update().get(pk=run.pk)
    if run.is_frozen:
        raise SamplingError("已冻结样本集不能继续导入候选论文；请创建新版本。")
    created = updated = 0
    for raw in records:
        if not isinstance(raw, dict):
            raise SamplingError("每条题录须为一个字段对象。")
        rec = normalize_record(raw, source_file)
        if not rec["title"]:
            raise SamplingError("题录缺少标题，请补全后重新导入。")
        if rec['url'] and urlparse(rec['url']).scheme not in ('http', 'https'):
            raise SamplingError('题录链接须为 HTTP 或 HTTPS 地址。')
        for field, length in [("title",1000),("journal",300),("issue",80),("volume",80),("doi",300),("url",1000),("cnki_filename",300),("cnki_dbcode",80),("clc",200),("article_type",200),("candidate_key",500)]:
            if len(str(rec.get(field, ""))) > length:
                raise SamplingError(f"题录字段 {field} 过长，请修正后重新导入。")
        fr = _attach_frame(run, rec)
        status, reason, decision = _classify(rec, fr)
        defaults = {
            "title": clean(rec.get("title")),
            "authors": clean(rec.get("authors")),
            "journal": clean(rec.get("journal")),
            "year": rec.get("year") or None,
            "issue": clean(rec.get("issue")),
            "volume": clean(rec.get("volume")),
            "doi": clean(rec.get("doi")),
            "cnki_url": clean(rec.get("url") or rec.get("cnki_url")),
            "cnki_dbcode": clean(rec.get("cnki_dbcode")),
            "cnki_filename": clean(rec.get("cnki_filename")),
            "keywords": clean(rec.get("keywords")),
            "abstract": clean(rec.get("abstract")),
            "clc": clean(rec.get("clc")),
            "article_type": clean(rec.get("article_type")),
            "stratum_id": fr["stratum_id"] if fr else "",
            "period": fr["period"] if fr else "",
            "tier": fr["tier"] if fr else "",
            "journal_family_name": fr["journal_family_name"] if fr else "",
            "frame_match_status": "MATCHED" if fr else "UNMATCHED",
            "cross_disciplinary_journal": bool(fr and fr['cross_disciplinary_journal']),
            "eligibility_status": status,
            "eligibility_reason": reason,
            "decision": decision,
            "source_file": (source_file or clean(rec.get("source_file")))[:300],
        }
        existing = run.candidates.filter(candidate_key=rec["candidate_key"]).first()
        if existing and existing.eligibility_status == CandidatePaper.UNCERTAIN and status == CandidatePaper.UNCERTAIN:
            defaults["decision"] = existing.decision
        _, was_created = CandidatePaper.objects.update_or_create(
            run=run, candidate_key=rec["candidate_key"], defaults=defaults
        )
        created += int(was_created)
        updated += int(not was_created)
    refresh_run_status(run)
    return {"created": created, "updated": updated, "total": run.candidates.count()}


def ingest_uploaded_file(run: SamplingRun, uploaded_file) -> dict:
    data = uploaded_file.read()
    try:
        records = read_metadata_bytes(uploaded_file.name, data)
    except Exception as exc:
        raise SamplingError("题录文件无法解析，请检查文件格式和内容。") from exc
    if not records:
        raise SamplingError("文件中没有识别到可用论文题录。")
    return ingest_records(run, records, uploaded_file.name)


def _issue_group(candidate: CandidatePaper):
    return f"{candidate.year or ''}-{candidate.issue or candidate.candidate_key[:8]}"


def _rank_candidates(candidates, seed: str, sid: str, method: str):
    rows = []
    for c in candidates:
        rows.append({
            "obj": c,
            "draw_hash": stable_hash(seed, sid, c.candidate_key),
            "issue_group": _issue_group(c),
        })
    if method == SamplingRun.METHOD_HASH_SIMPLE:
        rows.sort(key=lambda x: (x["draw_hash"], x["obj"].candidate_key))
    elif method == SamplingRun.METHOD_HASH_ISSUE_BALANCED:
        groups = defaultdict(list)
        for x in rows:
            groups[x["issue_group"]].append(x)
        for key in groups:
            groups[key].sort(key=lambda x: (x["draw_hash"], x["obj"].candidate_key))
        issue_order = sorted(groups, key=lambda key: (groups[key][0]["draw_hash"], key))
        ranked = []
        round_i = 0
        while True:
            added = False
            for key in issue_order:
                if round_i < len(groups[key]):
                    ranked.append(groups[key][round_i])
                    added = True
            if not added:
                break
            round_i += 1
        rows = ranked
    else:
        raise SamplingError("未知抽样方法。")
    for idx, x in enumerate(rows, 1):
        x["draw_rank"] = idx
    return rows


def _canonical_sha(rows):
    payload = json.dumps(rows, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return sha256_text(payload)


def _csv_bytes(fieldnames, rows):
    out = io.StringIO(newline="")
    w = csv.DictWriter(out, fieldnames=fieldnames, extrasaction="ignore")
    w.writeheader()
    for row in rows:
        w.writerow(row)
    return ("\ufeff" + out.getvalue()).encode("utf-8")


def _artifact_save(run, artifact_type, filename, data: bytes):
    old = list(run.artifacts.filter(artifact_type=artifact_type))
    for item in old:
        try:
            item.file.delete(save=False)
        except Exception:
            pass
        item.delete()
    obj = SamplingArtifact(
        run=run,
        artifact_type=artifact_type,
        original_name=filename,
        sha256=sha256_bytes(data),
    )
    obj.file.save(filename, ContentFile(data), save=True)
    return obj


def _candidate_registry_rows(run):
    fields = [
        "candidate_key", "title", "authors", "journal", "year", "issue", "volume",
        "doi", "cnki_url", "cnki_dbcode", "cnki_filename", "keywords", "abstract",
        "clc", "article_type", "stratum_id", "period", "tier",
        "journal_family_name", "frame_match_status", "cross_disciplinary_journal", "eligibility_status",
        "eligibility_reason", "decision",
    ]
    rows = []
    for c in run.candidates.order_by("candidate_key"):
        rows.append({f: getattr(c, f) for f in fields})
    return fields, rows


def _selected_rows(run):
    fields = [
        "paper_id", "sample_role", "reserve_for_role", "stratum_id", "tier", "period",
        "journal", "year", "issue", "title", "candidate_key", "draw_hash", "draw_rank",
        "stratum_pool_size", "reserve_rank",
    ]
    rows = []
    for c in run.candidates.exclude(sample_role="").order_by("paper_id"):
        rows.append({f: getattr(c, f) for f in fields})
    return fields, rows


def _build_download_queue(run):
    rows = []
    for c in run.candidates.filter(sample_role__in=["MAIN", "HOLDOUT"]).order_by("paper_id"):
        ids = parse_cnki(c.cnki_url, c.cnki_filename, c.cnki_dbcode)
        rows.append({
            "paper_id": c.paper_id, "role": c.sample_role, "tier": c.tier,
            "period": c.period, "journal": c.journal, "year": c.year or "",
            "issue": c.issue, "title": c.title, "cnki_url": c.cnki_url,
            "doi": c.doi, "cnki_record_key": ids["cnki_record_key"],
            "download_status": c.pdf_status,
        })
    return rows


def _build_evidence_queue(run):
    rows = []
    selected = run.candidates.filter(sample_role__in=["MAIN", "HOLDOUT"]).order_by("paper_id")
    for c in selected:
        ids = parse_cnki(c.cnki_url, c.cnki_filename, c.cnki_dbcode)
        required = 1 if ids["record_key_method"] in {"missing", "url_hash"} else 0
        rows.append({
            "task_id": f"{c.paper_id}-A",
            "paper_id": c.paper_id,
            "type": "论文题录",
            "required": required,
            "status": "待处理" if required else "自动信息已足够",
            "page": c.cnki_url or f"精确检索：{c.title}",
            "screenshot_name": f"{c.paper_id}_A_article_record.png" if required else "—",
            "what_to_capture": "题名 + 作者 + 期刊/来源 + 年份/期号；尽量保留地址栏" if required else "无需截图",
        })
    seen = set()
    i = 1
    for c in selected:
        if c.journal in seen:
            continue
        seen.add(c.journal)
        rows.append({
            "task_id": f"J{i:03d}-B",
            "paper_id": "—",
            "type": "期刊当前标签",
            "required": 1,
            "status": "待处理",
            "page": f"打开《{c.journal}》CNKI期刊/来源页",
            "screenshot_name": f"J{i:03d}_B_journal_tags.png",
            "what_to_capture": "期刊名称 + CNKI当前展示标签；同一期刊只截一次",
        })
        i += 1
    return rows


def _protocol_markdown(protocol: dict) -> bytes:
    text = f"""# 抽样协议与复现凭证

- 协议版本：`{protocol.get('protocol_version', '')}`
- 抽样方法：**{protocol.get('method_label', '')}** (`{protocol.get('method', '')}`)
- 主样本 seed：`{protocol.get('seed', '')}`
- 备用样本 seed：`{protocol.get('reserve_seed', '')}`
- Sampling frame SHA256：`{protocol.get('sampling_frame_sha256', '')}`
- Candidate registry SHA256：`{protocol.get('candidate_registry_sha256', '')}`
- Run fingerprint：`{protocol.get('run_fingerprint', '')}`

## 规则

{protocol.get('description', '')}

哈希优先级统一使用：

`SHA256(seed | stratum_id | candidate_key)`

相同抽样框、候选池、方法和 seed 应得到相同排序与样本。
"""
    return text.encode("utf-8")


@transaction.atomic
def freeze_run(run: SamplingRun):
    run = SamplingRun.objects.select_for_update().get(pk=run.pk)
    if run.is_frozen:
        raise SamplingError("该样本集已经冻结。")
    try:
        run.full_clean()
    except ValidationError as exc:
        raise SamplingError("请先完成研究范围设置：" + "；".join(exc.messages)) from exc
    unresolved = run.candidates.filter(
        eligibility_status=CandidatePaper.UNCERTAIN,
        decision=CandidatePaper.PENDING,
    ).count()
    if unresolved:
        raise SamplingError(f"仍有 {unresolved} 条待人工复核论文，不能冻结。")
    if not run.candidates.filter(decision=CandidatePaper.INCLUDE).exists():
        raise SamplingError("候选池中没有可抽取论文。")

    frame = build_frame(run)
    if not frame:
        raise SamplingError("请先完成研究范围设置，当前没有有效抽样格。")
    eligible = list(run.candidates.filter(decision=CandidatePaper.INCLUDE))
    by_stratum = defaultdict(list)
    for c in eligible:
        by_stratum[c.stratum_id].append(c)

    shortages = []
    for fr in frame:
        available = len(by_stratum.get(fr["stratum_id"], []))
        if available < int(fr["target_n"]):
            shortages.append({
                "stratum_id": fr["stratum_id"],
                "journal": fr["journal_family_name"],
                "period": fr["period"],
                "requested": int(fr["target_n"]),
                "available": available,
            })
    if shortages:
        raise SamplingShortage(shortages)

    run.candidates.update(
        sample_role="", paper_id="", draw_hash="", draw_rank=None, issue_group="",
        stratum_pool_size=None, reserve_rank=None, reserve_for_role=""
    )

    used = set()
    primary_counter = {"MAIN": 1, "HOLDOUT": 1}
    prefixes = {"MAIN": "P", "HOLDOUT": "H"}
    audit = []
    issues = []

    for role in ("MAIN", "HOLDOUT"):
        for fr in [x for x in frame if x["role"] == role]:
            sid = fr["stratum_id"]
            pool = [c for c in by_stratum.get(sid, []) if c.candidate_key not in used]
            ranked = _rank_candidates(pool, run.sampling_seed, sid, run.sampling_method)
            chosen = ranked[: int(fr["target_n"])]
            selected_keys = set()
            paper_ids = {}
            for row in chosen:
                c = row["obj"]
                c.sample_role = role
                c.paper_id = f"{prefixes[role]}{primary_counter[role]:03d}"
                c.draw_hash = row["draw_hash"]
                c.draw_rank = row["draw_rank"]
                c.issue_group = row["issue_group"]
                c.stratum_pool_size = len(ranked)
                c.reserve_rank = None
                c.reserve_for_role = ""
                c.save(update_fields=[
                    "sample_role", "paper_id", "draw_hash", "draw_rank",
                    "issue_group", "stratum_pool_size", "reserve_rank",
                    "reserve_for_role", "updated_at"
                ])
                primary_counter[role] += 1
                used.add(c.candidate_key)
                selected_keys.add(c.candidate_key)
                paper_ids[c.candidate_key] = c.paper_id
            for row in ranked:
                c = row["obj"]
                audit.append({
                    "protocol_version": PROTOCOL_VERSION,
                    "sampling_method": run.sampling_method,
                    "sampling_seed": run.sampling_seed,
                    "sampling_stage": f"{role}_DRAW",
                    "role": role,
                    "stratum_id": sid,
                    "pool_size": len(ranked),
                    "candidate_key": c.candidate_key,
                    "journal": c.journal,
                    "year": c.year or "",
                    "issue": c.issue,
                    "title": c.title,
                    "issue_group": row["issue_group"],
                    "draw_hash": row["draw_hash"],
                    "draw_rank": row["draw_rank"],
                    "selected": 1 if c.candidate_key in selected_keys else 0,
                    "paper_id": paper_ids.get(c.candidate_key, ""),
                })

    reserve_counter = 1
    reserve_seed = run.sampling_seed + "|reserve"
    for fr in frame:
        rn = int(fr["reserve_n"])
        if rn <= 0:
            continue
        sid = fr["stratum_id"]
        pool = [c for c in by_stratum.get(sid, []) if c.candidate_key not in used]
        ranked = _rank_candidates(pool, reserve_seed, sid, run.sampling_method)
        chosen = ranked[:rn]
        if len(chosen) < rn:
            issues.append({
                "stratum_id": sid,
                "role": fr["role"],
                "requested": rn,
                "selected": len(chosen),
                "issue": "备用样本不足",
            })
        selected_keys = set()
        paper_ids = {}
        for reserve_rank, row in enumerate(chosen, 1):
            c = row["obj"]
            c.sample_role = CandidatePaper.ROLE_RESERVE
            c.paper_id = f"R{reserve_counter:03d}"
            c.draw_hash = row["draw_hash"]
            c.draw_rank = row["draw_rank"]
            c.issue_group = row["issue_group"]
            c.stratum_pool_size = len(ranked)
            c.reserve_rank = reserve_rank
            c.reserve_for_role = fr["role"]
            c.save(update_fields=[
                "sample_role", "paper_id", "draw_hash", "draw_rank",
                "issue_group", "stratum_pool_size", "reserve_rank",
                "reserve_for_role", "updated_at"
            ])
            reserve_counter += 1
            used.add(c.candidate_key)
            selected_keys.add(c.candidate_key)
            paper_ids[c.candidate_key] = c.paper_id
        for row in ranked:
            c = row["obj"]
            audit.append({
                "protocol_version": PROTOCOL_VERSION,
                "sampling_method": run.sampling_method,
                "sampling_seed": reserve_seed,
                "sampling_stage": "RESERVE_DRAW",
                "role": fr["role"],
                "stratum_id": sid,
                "pool_size": len(ranked),
                "candidate_key": c.candidate_key,
                "journal": c.journal,
                "year": c.year or "",
                "issue": c.issue,
                "title": c.title,
                "issue_group": row["issue_group"],
                "draw_hash": row["draw_hash"],
                "draw_rank": row["draw_rank"],
                "selected": 1 if c.candidate_key in selected_keys else 0,
                "paper_id": paper_ids.get(c.candidate_key, ""),
            })

    frame_sha = _canonical_sha(frame)
    registry_fields, registry_rows = _candidate_registry_rows(run)
    registry_sha = _canonical_sha(registry_rows)
    method_label = dict(SamplingRun.METHOD_CHOICES).get(run.sampling_method, run.sampling_method)
    description = (
        "在每个期刊×时期×角色抽样格内计算 "
        "SHA256(seed | stratum_id | candidate_key)。"
    )
    if run.sampling_method == SamplingRun.METHOD_HASH_ISSUE_BALANCED:
        description += "先在各期号内部按哈希排序，再按期号轮转抽取，降低样本集中在同一期号的概率。"
    else:
        description += "按哈希值从小到大直接排序抽取。"
    fingerprint_payload = "|".join([
        PROTOCOL_VERSION, run.sampling_method, run.sampling_seed, frame_sha, registry_sha
    ])
    fingerprint = sha256_text(fingerprint_payload)
    protocol = {
        "protocol_version": PROTOCOL_VERSION,
        "method": run.sampling_method,
        "method_label": method_label,
        "seed": run.sampling_seed,
        "reserve_seed": reserve_seed,
        "description": description,
        "sampling_frame_sha256": frame_sha,
        "candidate_registry_sha256": registry_sha,
        "run_fingerprint": fingerprint,
        "frozen_at": timezone.now().isoformat(),
        "hash_format": "canonical_json_v1",
        "eligibility_rules_sha256": _canonical_sha(RULES),
        "journal_registry_sha256": _canonical_sha(REGISTRY),
        "frame_snapshot": frame,
    }

    run.protocol_version = PROTOCOL_VERSION
    run.sampling_frame_sha256 = frame_sha
    run.candidate_registry_sha256 = registry_sha
    run.run_fingerprint = fingerprint
    run.protocol = protocol
    run.status = SamplingRun.FROZEN
    run.frozen_at = timezone.now()
    run.save(update_fields=[
        "protocol_version", "sampling_frame_sha256", "candidate_registry_sha256",
        "run_fingerprint", "protocol", "status", "frozen_at", "updated_at"
    ])

    # Artifacts
    frame_fields = list(frame[0].keys()) if frame else []
    frame_csv = _csv_bytes(frame_fields, frame) if frame_fields else b""
    registry_csv = _csv_bytes(registry_fields, registry_rows)
    selected_fields, selected_rows = _selected_rows(run)
    selected_main = [x for x in selected_rows if x["sample_role"] in ("MAIN", "HOLDOUT")]
    selected_reserve = [x for x in selected_rows if x["sample_role"] == "RESERVE"]
    selected_csv = _csv_bytes(selected_fields, selected_main)
    reserve_csv = _csv_bytes(selected_fields, selected_reserve)

    audit_fields = [
        "protocol_version", "sampling_method", "sampling_seed", "sampling_stage",
        "role", "stratum_id", "pool_size", "candidate_key", "journal", "year",
        "issue", "title", "issue_group", "draw_hash", "draw_rank", "selected", "paper_id"
    ]
    ranking_csv = _csv_bytes(audit_fields, audit)

    certificate_rows = []
    for row in selected_rows:
        x = dict(row)
        x["run_fingerprint"] = fingerprint
        x["sampling_method"] = run.sampling_method
        x["sampling_seed"] = reserve_seed if row["sample_role"] == "RESERVE" else run.sampling_seed
        x["protocol_version"] = PROTOCOL_VERSION
        certificate_rows.append(x)
    certificate_fields = selected_fields + [
        "sampling_method", "sampling_seed", "protocol_version", "run_fingerprint"
    ]
    certificate_csv = _csv_bytes(certificate_fields, certificate_rows)
    issue_fields = ["stratum_id", "role", "requested", "selected", "issue"]
    issues_csv = _csv_bytes(issue_fields, issues)

    download_rows = _build_download_queue(run)
    download_fields = [
        "paper_id", "role", "tier", "period", "journal", "year", "issue",
        "title", "cnki_url", "doi", "cnki_record_key", "download_status",
    ]
    download_csv = _csv_bytes(download_fields, download_rows)
    evidence_rows = _build_evidence_queue(run)
    evidence_fields = [
        "task_id", "paper_id", "type", "required", "status", "page",
        "screenshot_name", "what_to_capture"
    ]
    evidence_csv = _csv_bytes(evidence_fields, evidence_rows)
    protocol_json = json.dumps(protocol, ensure_ascii=False, indent=2).encode("utf-8")
    protocol_md = _protocol_markdown(protocol)

    safe_fields = ["paper_id", "sample_role", "canonical_md", "pdf_status", "md_status"]
    safe_rows = [{"paper_id": c.paper_id, "sample_role": c.sample_role, "canonical_md": c.paper_id + ".md",
                  "pdf_status": c.pdf_status, "md_status": c.md_status} for c in run.candidates.exclude(paper_id="").order_by("paper_id")]
    safe_csv = _csv_bytes(safe_fields, safe_rows)
    label_fields = ["paper_id", "tier", "period", "journal", "year", "title", "authors", "cnki_url", "candidate_key"]
    label_rows = [{f: getattr(c, f) for f in label_fields} for c in run.candidates.exclude(paper_id="").order_by("paper_id")]
    labels_csv = _csv_bytes(label_fields, label_rows)
    files = {
        "sampling_frame.csv": frame_csv,
        "candidate_registry.csv": registry_csv,
        "selected_samples.csv": selected_csv,
        "reserve_samples.csv": reserve_csv,
        "sampling_issues.csv": issues_csv,
        "sampling_protocol.json": protocol_json,
        "SAMPLING_PROTOCOL.md": protocol_md,
        "sampling_certificate.csv": certificate_csv,
        "sampling_ranking_full.csv": ranking_csv,
        "download_queue.csv": download_csv,
        "evidence_queue.csv": evidence_csv,
        "safe_sample_manifest.csv": safe_csv,
        "labels_vault_seed_DO_NOT_COMMIT.csv": labels_csv,
    }
    bundle_io = io.BytesIO()
    with zipfile.ZipFile(bundle_io, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
    bundle = bundle_io.getvalue()

    _artifact_save(run, SamplingArtifact.TYPE_FRAME, "sampling_frame.csv", frame_csv)
    _artifact_save(run, SamplingArtifact.TYPE_REGISTRY, "candidate_registry.csv", registry_csv)
    _artifact_save(run, SamplingArtifact.TYPE_PROTOCOL, "sampling_protocol.json", protocol_json)
    _artifact_save(run, SamplingArtifact.TYPE_PROTOCOL_MD, "SAMPLING_PROTOCOL.md", protocol_md)
    _artifact_save(run, SamplingArtifact.TYPE_SAFE_MANIFEST, "safe_sample_manifest.csv", safe_csv)
    _artifact_save(run, SamplingArtifact.TYPE_LABELS, "labels_vault_seed_DO_NOT_COMMIT.csv", labels_csv)
    _artifact_save(run, SamplingArtifact.TYPE_CERTIFICATE, "sampling_certificate.csv", certificate_csv)
    _artifact_save(run, SamplingArtifact.TYPE_RANKING, "sampling_ranking_full.csv", ranking_csv)
    _artifact_save(run, SamplingArtifact.TYPE_SELECTED, "selected_samples.csv", selected_csv)
    _artifact_save(run, SamplingArtifact.TYPE_RESERVE, "reserve_samples.csv", reserve_csv)
    _artifact_save(run, SamplingArtifact.TYPE_ISSUES, "sampling_issues.csv", issues_csv)
    _artifact_save(run, SamplingArtifact.TYPE_DOWNLOAD_QUEUE, "download_queue.csv", download_csv)
    _artifact_save(run, SamplingArtifact.TYPE_EVIDENCE_QUEUE, "evidence_queue.csv", evidence_csv)
    _artifact_save(run, SamplingArtifact.TYPE_BUNDLE, "sampling_result_bundle.zip", bundle)

    return {"protocol": protocol, "issues": issues}


@transaction.atomic
def update_review_decisions(run: SamplingRun, ids: list[int], action: str):
    run = SamplingRun.objects.select_for_update().get(pk=run.pk)
    if run.is_frozen:
        raise SamplingError("已冻结样本集不能修改复核结果。")
    if action not in (CandidatePaper.INCLUDE, CandidatePaper.EXCLUDE):
        raise SamplingError("未知复核操作。")
    qs = run.candidates.filter(
        pk__in=ids, eligibility_status=CandidatePaper.UNCERTAIN
    )
    count = qs.update(decision=action, updated_at=timezone.now())
    refresh_run_status(run)
    return count


def _read_csv_from_zip(z, name):
    try:
        data = z.read(name)
    except KeyError:
        return []
    return list(csv.DictReader(io.StringIO(_decode_bytes(data))))


def import_result_bundle(project, uploaded_file, user, name="", version="imported"):
    from .bundles import import_bundle
    return import_bundle(project, uploaded_file, user, name, version)


def collection_config(run: SamplingRun):
    return {
        "run_id": run.pk,
        "scope_signature": _canonical_sha({"frame": build_frame(run), "method": run.sampling_method, "seed": run.sampling_seed}),
        "periods": run.periods,
        "selected_tiers": run.selected_tiers,
        "selected_journals": run.selected_journals,
        "main_n": run.main_n,
        "reserve_n": run.reserve_n,
        "holdout_enabled": run.holdout_enabled,
        "holdout_start": run.holdout_start,
        "holdout_end": run.holdout_end,
        "holdout_n": run.holdout_n,
        "frame": build_frame(run),
    }


def freeze_readiness(run):
    errors = validate_scope(run)
    if errors:
        return {'ready': False, 'message': '请先完成研究范围设置', 'shortages': []}
    if run.is_frozen:
        return {'ready': False, 'message': '该版本已冻结', 'shortages': []}
    unresolved = run.candidates.filter(eligibility_status=CandidatePaper.UNCERTAIN, decision='').count()
    if unresolved:
        return {'ready': False, 'message': f'请先完成 {unresolved} 篇论文的入选复核', 'shortages': []}
    pools = dict(run.candidates.filter(decision='include').values_list('stratum_id').annotate(total=Count('pk')))
    frame = build_frame(run)
    shortages = [dict(row, available=pools.get(row['stratum_id'], 0)) for row in frame
                 if pools.get(row['stratum_id'], 0) < row['target_n']]
    if not frame or not pools:
        return {'ready': False, 'message': '请先采集或导入候选题录', 'shortages': shortages}
    return {'ready': not shortages, 'message': f'有 {len(shortages)} 个抽样格主样本不足' if shortages else '研究范围完整，候选池已满足冻结条件', 'shortages': shortages}
